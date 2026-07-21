from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


OUT_PATH = Path(r"C:\Users\ziwei.luo\Documents\New project\outputs\20260720\shouban_guide_sheet2_filled.xlsx")

SHEET2_ROWS = [
    (
        "UIGKCollectionRoom",
        "收藏室主页",
        "收藏室主界面 prefab，包含 15 个展示位、顶栏、编辑模式以及 Add/Remove/Exchange 相关操作。"
        "逻辑节点包括 EditBtn、SaveBtn、QuitBtn、ResetBtn、EditMode、LeftCollectionRoot、"
        "ModelIconEastList、ActivityPopMenuTrans/SortPopMenuTrans、ModelNumTrans/SelectModelLabel、"
        "GiftBtn/ExchangeBtn/EmoteBtn/PreviewModelBtn、ValueCntLabel/ProgressCntLabel/ActivityLabel、"
        "ERankBtn/ShareBtn/FriendListBtn。筛选/排序弹出菜单挂载 UIPopMenuSmall，底部选中手办编号显示"
        "挂到 ModelNumTrans 的 UIGKModelNum。"
    ),
    (
        "UIGKModelSlot",
        "手办柜单格 / 展示柜格子",
        "单个展示柜格子 prefab，用于 3D 模型挂点和增删换按钮状态机。大小通过根节点 scale 区分，"
        "需要同步大柜子 scale。关键逻辑节点包括 MainBtn 以及 AddBtn/ExchangeBtn/RemoveBtn；"
        "MainBtn 负责选中、放置、换位路由，相关 surfaceMain 节点需挂 BoxCollider 作为拖拽落点保障，"
        "按钮显隐通过 RefreshButtons 联动。"
    ),
    (
        "UIGKModelIconItem",
        "背包 / 列表格子 / 背包单项",
        "手办 Icon 项 prefab，用于背包或列表单格展示。包含 ItemBtn、DragItem+DragItemIcon、"
        "itemSprite/ItemIconView2、HL、Quality/QualityBG、Cnt、SpecialNumTrans 等节点。"
        "支持选中/取消选中、运行时拖拽克隆、稀有度框显示、数量显示以及编号挂点；"
        "SpecialNumTrans 挂 UIGKModelNum，Number 大于 0 时显示。"
    ),
    (
        "UIGKModelNum",
        "手办编号显示组件 / 编号锚点",
        "手办编号显示组件，其他 prefab 的 numTrans 节点需要生成并挂载此 prefab。包含 Num/NumBG 文本、"
        "UIGKModelNum 面板以及 DetailTipsCamera。编号带独立相机渲染，重新导出时不能丢失；"
        "主要用于底部选中手办信息、背包单项编号和汇总弹窗编号显示。"
    ),
    (
        "UIGKCatalogPopup",
        "图鉴弹窗",
        "图鉴弹窗 prefab。主体节点包括 TabsContainer、GKModelTable、GKCatalogScrollView、"
        "ActivityLabel、CloseBtn。GKModelTable 用于混排 section title 和 catalog item，"
        "GKCatalogScrollView 为图鉴滚动区域，CloseBtn 用于关闭图鉴；"
        "ActivityLabel 当前导出但 controller 未实际设置。"
    ),
    (
        "UIGKCatalogItem",
        "图鉴网格单格 / 图鉴单项",
        "图鉴单项 prefab，用于图鉴网格中的单格展示。主要节点包括 ItemIcon、QualityBG、Locked。"
        "ItemIcon 根据 IconResId 解析图标，未解锁时置灰；QualityBG 控制稀有度着色，"
        "Locked 用于未解锁遮罩并受 IsUnlocked 控制。"
    ),
    (
        "UIGKCatalogSectionTitle",
        "图鉴品质标题",
        "图鉴品质分段标题 prefab，用于图鉴列表中不同品质区块的标题展示。PDF 中未给出更多节点拆分，"
        "当前可确认其用途为图鉴分组标题显示。"
    ),
    (
        "UIGKCatalogTab",
        "图鉴系列 Tab 容器",
        "图鉴系列 Tab 容器 prefab，用于承载图鉴系列切页项。PDF 中未给出额外节点说明，"
        "当前可确认其职责为系列 Tab 的容器与布局承载。"
    ),
    (
        "UIGKCatalogTabItem",
        "图鉴系列 Tab 单项",
        "图鉴系列 Tab 单项 prefab。PDF 提取文本中显示为 UIGKCatalogTabtem，结合语义按 TabItem 归一。"
        "当前可确认其用于图鉴系列切页的单个 Tab 项展示与交互。"
    ),
    (
        "UIGKCollectionNewGainPopWindow",
        "新获得手办弹窗（掉落 / 解锁提示）",
        "新获得手办的拍脸弹窗 prefab。关键内容节点包括 ModelSprite、GKModelName、"
        "SourceTimeAndMapLabel、NumTrans，分别用于显示手办图、名称、带出时间与地图、编号锚点。"
        "数据依赖 OB56 带出的结算协议。"
    ),
    (
        "UIGKCollectionNewGainAllPopWindow",
        "汇总弹窗",
        "手办获取汇总弹窗 prefab。PDF 中当前仅明确为汇总弹窗容器，详细节点拆解未单独列出；"
        "通常与 UIGKNewGainItem 配合显示批量获取结果。"
    ),
    (
        "UIGKNewGainItem",
        "汇总弹窗 Item",
        "汇总弹窗中的单个 Item prefab。当前作用节点包括 ItemButton、SpecialNumTrans、FirstTag、"
        "QualityLine、QualityBG、itemSprite、ValueLabel。支持点击打开图鉴 tips、编号显示、"
        "首次获得标记以及品质/图标/收藏值展示。"
    ),
    (
        "UIGKCollectionGKCatalogTipsBox",
        "大厅图鉴解锁提示",
        "大厅图鉴解锁相关 tips/popup prefab。PDF 中仅明确名称与用途为大厅图鉴解锁提示，"
        "未给出进一步节点说明。"
    ),
    (
        "UILobbyV2GameAssistNewGKNoticeItem",
        "小助手提示 Item",
        "大厅小助手中的手办提示 Item prefab。PDF 中仅明确其为小助手提示项，"
        "适用于相关引导或新内容提示展示。"
    ),
    (
        "UIGKCatalogItemTips",
        "Item 点击 tips",
        "图鉴 item 点击后展开的 tips prefab。关键节点包括 MaskBtn、ContentTable、SmallItemTrans、"
        "FirstGainLabel、TimeLabel、MapAndPoiLabel、TeammateGrid。支持点击遮罩关闭、tips 内容重排、"
        "小手办 icon 子 controller 挂点、获得信息区域整体显隐、首次获得时间、地图/地点与队友来源展示。"
    ),
    (
        "UIGKCatalogTipsTeammate",
        "tips 里的队友条目",
        "图鉴 tips 内的队友来源条目 prefab，用于展示局内获得队友列表，或赠礼 / 交换来源玩家信息。"
        "PDF 未进一步拆解节点。"
    ),
    (
        "BaseHeadTrans",
        "挂载 UIBaseHead",
        "头像挂点节点，用于挂载 UIBaseHead。PDF 将其列为 prefab 相关项，但更像是约定挂点或基础头部容器，"
        "通常与队友或来源展示配合。"
    ),
    (
        "UIGKSmallItem",
        "小助手 / tips 里的小尺寸 item",
        "小尺寸手办 item prefab，供小助手和 tips 内部复用。适合在信息密度较高的提示区域展示缩略手办信息。"
    ),
    (
        "UIGKEmoteUnLockPopWindow",
        "Emote 获取弹窗",
        "Emote 解锁 / 获取弹窗 prefab。PDF 中仅明确用途为 Emote 获取提示，未列出更多节点结构。"
    ),
    (
        "UIGKEmoteChoose",
        "手办展示界面内 item 复用",
        "手办展示界面中的选择项 prefab，内部 item 复用 UIGKModelIconItem。"
        "用于在手办展示或 Emote 相关界面中复用同套 icon 交互展示。"
    ),
    (
        "UINewVault",
        "新增更换 GK 按钮",
        "新增更换 GK 按钮相关 prefab，PDF 中关联文案为 ChangeGKBtn。"
        "当前可确认其职责为更换 GK 的入口按钮或容器。"
    ),
]


def style_sheet(ws, widths):
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    body_font = Font(name="Microsoft YaHei", size=11)
    header_font = Font(name="Microsoft YaHei", size=11, bold=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = left

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1_UI指路总表"
    ws2 = wb.create_sheet("Sheet2_Prefab说明")
    ws3 = wb.create_sheet("Sheet3_CDN配置")
    ws4 = wb.create_sheet("Sheet4_需要注册的图片")

    ws1.append(["功能名称", "值"])
    for row in [
        ("功能名称", "手办系统最终"),
        ("设计人", "待补充"),
        ("制作人", "待补充"),
        ("策划", "待补充"),
        ("UE", "待补充"),
        ("程序", "待补充"),
        ("Jira链接", "待补充"),
        ("Wiki链接", "待补充"),
        ("图集名称", "待补充"),
    ]:
        ws1.append(list(row))
    style_sheet(ws1, {"A": 16, "B": 60})

    ws2.append(["Prefab名称", "位置", "Prefab详细说明"])
    for row in SHEET2_ROWS:
        ws2.append(list(row))
    style_sheet(ws2, {"A": 34, "B": 30, "C": 92})
    for idx in range(1, ws2.max_row + 1):
        ws2.row_dimensions[idx].height = 48

    ws3.append(["图片类目", "CDN位置描述", "配置名称", "CDN路径"])
    ws3.append(["待补充", "待补充", "待补充", "待补充"])
    style_sheet(ws3, {"A": 20, "B": 28, "C": 28, "D": 60})

    ws4.append(["截图", "图集名称", "图片名称"])
    ws4.append(["插入截图", "待补充", "待补充"])
    style_sheet(ws4, {"A": 24, "B": 26, "C": 34})

    wb.save(OUT_PATH)

    verify = load_workbook(OUT_PATH)
    sheet2 = verify["Sheet2_Prefab说明"]
    print(OUT_PATH)
    print(f"sheet2_rows={sheet2.max_row}")
    print(f"first_prefab={sheet2['A2'].value}")
    print(f"last_prefab={sheet2[f'A{sheet2.max_row}'].value}")


if __name__ == "__main__":
    main()
